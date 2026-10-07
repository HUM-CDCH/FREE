"""The developer iterative-eval harness: golden Excel, field-level P/R/F1, pilot->full, improvement."""
from __future__ import annotations

import csv
import json
from pathlib import Path

import openpyxl
import pytest

from experiments.extraction import iterative_eval

SCHEMA = {
    "recordDescription": "One dated sample.",
    "recordScope": "document",
    "schemaNodes": [
        {"id": "n", "name": "sample_name", "type": "string"},
        {"id": "y", "name": "year", "type": "integer"},
        {"id": "m", "name": "materials", "type": "array",
         "children": [{"id": "mi", "name": "item", "type": "string"}]},
    ],
}


def _golden(tmp_path: Path) -> Path:
    path = tmp_path / "golden.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "gold"
    sheet.append(["file", "sample_name", "year", "materials"])
    sheet.append(["sources/paper-a.pdf", "Alpha", 1990, '["bone", "antler"]'])
    sheet.append(["paper-b", "Beta", 2000, "bone"])  # blank-free single value; second row names the file plainly
    sheet.append([None, None, None, None])
    workbook.save(path)
    return path


def test_load_gold_reads_columns_and_cells(tmp_path: Path):
    gold = iterative_eval.load_gold(_golden(tmp_path))
    assert gold["document_column"] == "file"
    assert gold["fields"] == ["sample_name", "year", "materials"]
    assert set(gold["documents"]) == {"paper-a", "paper-b"}
    assert gold["documents"]["paper-a"][0]["materials"] == ["bone", "antler"]
    assert gold["documents"]["paper-b"][0]["year"] == [2000]
    assert gold["documents"]["paper-b"][0]["materials"] == ["bone"]


def test_score_counts_precision_recall_and_f1(tmp_path: Path):
    gold = iterative_eval.load_gold(_golden(tmp_path))
    predictions = {
        "paper-a": {"records": [{"sample_name": "alpha", "year": 1990, "materials": ["bone", "antler"]}]},
        "paper-b": {"records": [{"sample_name": "Beta", "year": 2005, "materials": ["bone", "shell"]}]},
    }
    metrics = iterative_eval.score(gold, predictions, phase="full", eval_id="t")
    assert metrics["micro"]["tp"] == 6  # sample_name x2, year x1, materials x3
    assert metrics["micro"]["fp"] == 2  # the extra 2005 and the extra "shell"
    assert metrics["micro"]["fn"] == 1  # the missed 2000
    assert metrics["micro"]["precision"] == pytest.approx(6 / 8)
    assert metrics["micro"]["recall"] == pytest.approx(6 / 7)
    assert metrics["presence"]["f1"] == 1.0  # every field was at least populated
    assert metrics["per_field"]["year"]["fp"] == 1 and metrics["per_field"]["year"]["fn"] == 1
    assert metrics["exact_cells"] == {"matched": 4, "total": 6, "accuracy": pytest.approx(4 / 6)}
    assert "| sample_name |" in iterative_eval.render(metrics)


def test_compare_reports_the_improvement(tmp_path: Path):
    gold = iterative_eval.load_gold(_golden(tmp_path))
    weak = iterative_eval.score(gold, {"paper-a": None, "paper-b": None}, phase="pilot", eval_id="v1")
    strong = iterative_eval.score(gold, {"paper-a": {"records": [{"sample_name": "Alpha", "year": 1990}]},
                                         "paper-b": {"records": [{"sample_name": "Beta", "year": 2000}]}},
                                  phase="pilot", eval_id="v2")
    report = iterative_eval.compare(weak, strong)
    assert report["delta"]["micro"]["f1"] > 0
    assert "improved" in report["conclusion"]
    assert "sample_name" in report["per_field"]


def _run_dir(root: Path, key: str) -> Path:
    directory = root / "runs" / key
    directory.mkdir(parents=True)
    (directory / "params.json").write_text(json.dumps({"source_name": f"{key}.pdf"}), encoding="utf-8")
    return directory


def test_run_eval_pilots_then_covers_the_whole_set_and_reuses_cells(tmp_path: Path, monkeypatch):
    (tmp_path / "schema.json").write_text(json.dumps(SCHEMA), encoding="utf-8")
    _golden(tmp_path)
    config = {
        "id": "t", "schema": "schema.json", "golden": "golden.xlsx", "output": "eval-out", "pilot": 1,
        "options": {"strategy": "article"},
        "providers": {"fields": {"base_url": "http://x", "model": "f"},
                      "reasoning": {"base_url": "http://x", "model": "r"}},
        "documents": ["paper-a.pdf", "paper-b.pdf"],
        "runs": {"paper-a": "runs/paper-a", "paper-b": "runs/paper-b"},
    }
    _run_dir(tmp_path, "paper-a")
    _run_dir(tmp_path, "paper-b")
    artifacts = {
        "paper-a": {"records": [{"sample_name": "Alpha", "year": 1990, "materials": ["bone", "antler"]}]},
        "paper-b": {"records": [{"sample_name": "Beta", "year": 2000, "materials": ["bone"]}]},
    }
    calls = []

    def fake(run_dir, request, providers, cell_dir):
        calls.append(run_dir.name)
        return artifacts[run_dir.name]

    monkeypatch.setattr(iterative_eval, "_extract_document", fake)
    pilot, full = iterative_eval.run_eval(config, tmp_path, phase="full")
    assert pilot["documents"] == ["paper-a"] and pilot["micro"]["f1"] == 1.0
    assert full["documents"] == ["paper-a", "paper-b"] and full["micro"]["f1"] == 1.0
    assert (tmp_path / "eval-out" / "metrics-full.json").is_file()
    assert (tmp_path / "eval-out" / "report-pilot.md").is_file()
    assert len(calls) == 2
    iterative_eval.run_eval(config, tmp_path, phase="full")  # completed cells are reused, not re-extracted
    assert len(calls) == 2


def _golden_three(tmp_path: Path) -> Path:
    path = tmp_path / "golden-three.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "gold"
    sheet.append(["file", "sample_name", "year", "materials"])
    sheet.append(["paper-a", "Alpha", 1990, '["bone"]'])
    sheet.append(["paper-b", "Beta", 2000, "bone"])
    sheet.append(["paper-c", "Gamma", 2010, "shell"])
    workbook.save(path)
    return path


def test_guidance_examples_list_only_differing_fields(tmp_path: Path):
    gold = iterative_eval.load_gold(_golden(tmp_path))
    matching = {
        "paper-a": {
            "records": [
                {"sample_name": "alpha", "year": 1990, "materials": ["bone", "antler"]}
            ]
        }
    }
    assert iterative_eval.guidance_examples(gold, matching, ["paper-a"]) == []
    assert iterative_eval.guidance_text([]) == ""

    differing = {
        "paper-a": {
            "records": [
                {"sample_name": "Alpha", "year": 1985, "materials": ["bone", "antler"]}
            ]
        }
    }
    examples = iterative_eval.guidance_examples(gold, differing, ["paper-a"])
    assert examples == [{"field": "year", "expected": 1990}]
    assert iterative_eval.guidance_text(examples).endswith(
        json.dumps(examples, ensure_ascii=False, sort_keys=True)
    )


def test_guidance_counter_counts_the_block_the_chat_will_send():
    class FakeCounter:
        context_tokens = 1000

        def __init__(self):
            self.seen: list[str] = []

        def request_tokens(self, system, user, schema=None):
            self.seen.append(system)
            return len(system)

    counter = FakeCounter()
    guided = iterative_eval.GuidanceCounter(counter, "\nEXTRA")
    assert guided.request_tokens("sys", "user") == len("sys\nEXTRA")
    assert guided.context_tokens == 1000
    assert counter.seen == ["sys\nEXTRA"]


def test_provider_client_kwargs_bounds_only_the_fields_role():
    assert iterative_eval.provider_client_kwargs("fields", 16) == {"max_whitespace": 16}
    assert iterative_eval.provider_client_kwargs("reasoning", 16) == {}
    assert iterative_eval.provider_client_kwargs("fields", 0) == {}
    assert iterative_eval.DEFAULT_MAX_WHITESPACE == 16


class _FakeReply:
    def __init__(self, text: str, finish: str = "stop"):
        self.text, self.finish = text, finish


class _FakeCapture:
    def __init__(self, replies: list):
        self.replies, self.requests = list(replies), []

    def complete(self, **request):
        self.requests.append(request)
        return self.replies.pop(0)


def test_judge_pair_maps_a_verdict_to_counts():
    payload = {
        "extracted": [
            {"value": "a", "verdict": "match"},
            {"value": "b", "verdict": "extra"},
            {"value": "c", "verdict": "uncertain"},
        ],
        "missing_gold": ["d"],
    }
    result = iterative_eval.judge_pair(
        _FakeCapture([_FakeReply(json.dumps(payload))]), "f", ["a"], ["a", "b", "c"]
    )
    assert result == {"tp": 1, "fp": 1, "uncertain": 1, "fn": 1, "verdict": payload}


def test_judge_pair_leaves_a_failed_judge_unjudged():
    assert iterative_eval.judge_pair(_FakeCapture([_FakeReply("{}", finish="length")]), "f", ["a"], ["b"]) is None
    assert iterative_eval.judge_pair(_FakeCapture([_FakeReply("not json")]), "f", ["a"], ["b"]) is None


def test_judge_round_skips_pairs_strict_matching_already_confirmed():
    gold = {"fields": ["f"], "documents": {"a": [{"f": ["x"]}]}}
    predictions = {"a": {"records": [{"f": "x"}]}}
    capture = _FakeCapture([])  # a judge call would raise IndexError
    result = iterative_eval.judge_round(gold, predictions, ["a"], capture, {}, model="m")
    assert capture.requests == []
    assert result["judged"] == 1 and result["unjudged"] == 0
    assert result["micro"]["tp"] == 1 and result["micro"]["f1"] == 1.0


def test_guidance_text_caps_the_examples_it_uses():
    examples = [{"field": "f", "expected": "x" * 4000} for _ in range(5)]
    capped = iterative_eval.capped_examples(examples, max_chars=8000)
    assert 0 < len(capped) < len(examples)
    assert len(iterative_eval.guidance_text(examples, max_chars=8000)) <= 8000


def test_configured_pilot_entries_defaults_to_the_first_two(tmp_path: Path):
    entries = [{"key": key, "pdf": None, "run": None} for key in ("a", "b", "c")]
    assert [entry["key"] for entry in iterative_eval.configured_pilot_entries({}, entries)] == ["a", "b"]
    assert [
        entry["key"]
        for entry in iterative_eval.configured_pilot_entries({"pilot_documents": ["c", "a"]}, entries)
    ] == ["c", "a"]
    with pytest.raises(ValueError):
        iterative_eval.configured_pilot_entries({}, entries[:1])


def test_anchor_coverage_is_the_share_of_eligible_populated_leaves():
    artifact = {
        "schema": {"schemaNodes": [{"id": "n", "name": "sample_name"}]},
        "records": [{"sample_name": "Alpha", "year": 1990}],
        "evidence": [{"path": ["records", 0, "sample_name"]}],
        "unverified": [],
        "grounding_eligibility": {
            "all_record_leaves": 2,
            "eligible_record_leaves": 1,
            "skipped": [{"path": ["records", 0, "year"], "policy": "derived"}],
        },
    }
    coverage = iterative_eval.anchor_coverage({"paper-a": artifact})
    assert coverage["populated_record_leaves"] == 2
    assert coverage["linked_record_leaves"] == 1
    assert coverage["link_rate"] == pytest.approx(0.5)
    assert coverage["eligible_record_leaves"] == 1
    assert coverage["linked_eligible_record_leaves"] == 1
    assert coverage["eligible_link_rate"] == pytest.approx(1.0)
    assert coverage["skipped_record_leaves"] == 1
    assert "not semantic correctness" in coverage["interpretation"]


def test_shadow_effort_classifies_edits_rejections_additions_and_deletions():
    gold = {
        "fields": ["f"],
        "documents": {
            "a": [{"f": ["keep", "drop"]}],
            "b": [{"f": []}],
            "c": [{"f": ["new"]}],
            "d": [{"f": ["only-extra"]}],
        },
    }
    predictions = {
        "a": {"records": [{"f": ["keep", "extra"]}]},        # a replaced value -> one field edit
        "b": {"records": [{"f": "wrong"}]},                  # gold is empty -> rejection
        "c": {"records": [{"f": []}]},                        # nothing extracted -> addition
        "d": {"records": [{"f": ["only-extra", "more"]}]},   # one extra value -> deletion
    }
    effort = iterative_eval.shadow_effort(gold, predictions, ["a", "b", "c", "d"])
    assert {key: effort[key] for key in ("edited", "rejected", "added", "deleted", "effort")} == {
        "edited": 1,
        "rejected": 1,
        "added": 1,
        "deleted": 1,
        "effort": 4,
    }
    assert "no decision" in effort["interpretation"]


def test_score_treats_extra_predictions_by_gold_exhaustiveness(tmp_path: Path):
    gold = iterative_eval.load_gold(_golden(tmp_path))
    predictions = {
        "paper-a": {
            "records": [
                {"sample_name": "Alpha", "year": 1990, "materials": ["bone", "antler", "extra"]}
            ]
        },
        "paper-b": {"records": [{"sample_name": "Beta", "year": 2000, "materials": ["bone"]}]},
    }

    exhaustive = iterative_eval.score(gold, predictions, phase="p", eval_id="t")
    assert exhaustive["micro"]["tp"] == 7
    assert exhaustive["micro"]["fp"] == 1  # the extra material, the default standard-answer convention
    assert exhaustive["micro"]["precision_computed"] is True
    assert exhaustive["unscored_extras"]["count"] == 0

    partial = iterative_eval.score(gold, predictions, phase="p", eval_id="t",
                                   options={"exhaustive": False})
    assert partial["micro"]["tp"] == 7
    assert partial["micro"]["fp"] == 0
    assert partial["micro"]["precision"] is None
    assert partial["micro"]["recall"] == pytest.approx(1.0)
    assert partial["micro"]["f1"] is None
    assert partial["unscored_extras"] == {
        "count": 1,
        "exhaustive": False,
        "by_field": {"sample_name": 0, "year": 0, "materials": 1},
    }
    assert "precision/F1 not computed" in iterative_eval.render(partial)


def test_align_records_matches_by_identity_and_reports_unmatched():
    gold_rows = [
        {"amino_acid_hydroxyproline_value": ["1"], "species": ["Bone"]},
        {"amino_acid_hydroxyproline_value": ["2"], "species": ["Antler"]},
    ]
    records = [
        {"amino_acid_hydroxyproline_value": "2", "species": "antler"},
        {"amino_acid_hydroxyproline_value": "3", "species": "Shell"},
        {"species": "No key"},
    ]
    alignment = iterative_eval.align_records(
        gold_rows, records, ["amino_acid_hydroxyproline_value"], {}
    )
    assert [(row["species"], record["species"]) for row, record in alignment["pairs"]] == [
        (["Antler"], "antler")
    ]
    assert [row["species"] for row in alignment["unmatched_gold_rows"]] == [["Bone"]]
    assert [record["species"] for record in alignment["unmatched_predicted_records"]] == [
        "Shell",
        "No key",
    ]


def test_align_records_leaves_an_ambiguous_identity_unmatched():
    gold_rows = [{"k": ["1"]}, {"k": ["1"]}]
    records = [{"k": "1"}]
    alignment = iterative_eval.align_records(gold_rows, records, ["k"], {})
    assert alignment["pairs"] == []
    assert len(alignment["unmatched_gold_rows"]) == 2
    assert len(alignment["unmatched_predicted_records"]) == 1


def test_score_aligns_catalog_records_by_the_identity_field():
    gold = {
        "path": "gold",
        "sha256": "x",
        "fields": ["amino_acid_hydroxyproline_value", "species"],
        "documents": {
            "paper": [
                {"amino_acid_hydroxyproline_value": ["1"], "species": ["Bone"]},
                {"amino_acid_hydroxyproline_value": ["2"], "species": ["Antler"]},
            ]
        },
    }
    predictions = {
        "paper": {
            "records": [
                {"amino_acid_hydroxyproline_value": "2", "species": "Antler"},
                {"amino_acid_hydroxyproline_value": "3", "species": "Shell"},
            ]
        }
    }
    metrics = iterative_eval.score(gold, predictions, phase="p", eval_id="t")
    assert metrics["record_alignment"] == {
        "aligned": 1,
        "unmatched_gold": 1,
        "unmatched_predicted": 1,
    }
    assert metrics["micro"]["tp"] == 2
    assert metrics["micro"]["fp"] == 2
    assert metrics["micro"]["fn"] == 2

    partial = iterative_eval.score(
        gold, predictions, phase="p", eval_id="t", options={"exhaustive": False}
    )
    assert partial["micro"]["tp"] == 2
    assert partial["micro"]["fp"] == 0
    assert partial["micro"]["precision"] is None
    assert partial["unscored_extras"]["count"] == 2


def test_metrics_row_is_one_row_per_round() -> None:
    metrics = {
        "micro": {"precision": 0.5, "recall": 0.6, "f1": 0.54},
        "presence": {"precision": 1.0, "recall": 0.5, "f1": 0.67},
        "exact_cells": {"accuracy": 0.25},
        "documents_scored": 2,
        "documents_failed": ["c"],
        "anchor_coverage": {"eligible_link_rate": 0.9, "link_rate": 0.8},
        "reviewer_effort": {"effort": 4, "edited": 1, "rejected": 1, "added": 1, "deleted": 1},
    }
    row = iterative_eval.metrics_row(
        "PILOT_1",
        {"status": "SUCCEEDED", "documents": ["a", "b"], "guidance_sha256": "g"},
        metrics,
        gold_sha256="gold",
    )
    assert row["Round"] == "PILOT_1"
    assert row["F1"] == 0.54
    assert row["Anchor coverage"] == 0.9
    assert row["Effort"] == 4
    assert row["Documents"] == 2 and row["Scored"] == 2 and row["Failed"] == 1


def test_write_metrics_table_writes_csv_and_xlsx(tmp_path: Path) -> None:
    rows = [
        iterative_eval.metrics_row("PILOT_1", {"status": "SUCCEEDED", "documents": []}, None,
                                   gold_sha256="g")
    ]
    iterative_eval.write_metrics_table(tmp_path, rows)
    assert (tmp_path / "metrics.csv").is_file()
    sheet = openpyxl.load_workbook(tmp_path / "metrics.xlsx").active
    assert [cell.value for cell in sheet[1]][:2] == ["Round", "Status"]
    assert sheet[2][0].value == "PILOT_1"


def test_run_pipeline_runs_two_pilots_then_a_batch_with_guidance(tmp_path: Path, monkeypatch):
    (tmp_path / "schema.json").write_text(json.dumps(SCHEMA), encoding="utf-8")
    _golden_three(tmp_path)
    for key in ("paper-a", "paper-b", "paper-c"):
        _run_dir(tmp_path, key)
    config = {
        "id": "p",
        "schema": "schema.json",
        "golden": "golden-three.xlsx",
        "output": "pipe",
        "options": {"strategy": "article"},
        "providers": {
            "fields": {"base_url": "http://x", "model": "f"},
            "reasoning": {"base_url": "http://x", "model": "r"},
        },
        "documents": ["paper-a.pdf", "paper-b.pdf", "paper-c.pdf"],
        "runs": {
            "paper-a": "runs/paper-a",
            "paper-b": "runs/paper-b",
            "paper-c": "runs/paper-c",
        },
    }
    artifacts = {
        # paper-a misses the gold year, so pilot-2 and the batch receive guidance.
        "paper-a": {"records": [{"sample_name": "Alpha", "year": 1985, "materials": ["bone"]}]},
        "paper-b": {"records": [{"sample_name": "Beta", "year": 2000, "materials": ["bone"]}]},
        "paper-c": {"records": [{"sample_name": "Gamma", "year": 2010, "materials": ["shell"]}]},
    }
    calls: list[tuple[str, str]] = []

    def fake(run_dir, request, providers, cell_dir, *, guidance="", max_whitespace=16):
        calls.append((run_dir.name, guidance))
        return artifacts[run_dir.name]

    monkeypatch.setattr(iterative_eval, "_extract_document", fake)
    result = iterative_eval.run_pipeline(config, tmp_path)

    assert [round["label"] for round in result["rounds"]] == ["PILOT_1", "PILOT_2", "BATCH"]
    assert [round["status"] for round in result["rounds"]] == ["SUCCEEDED", "SUCCEEDED", "SUCCEEDED"]
    assert [name for name, _ in calls] == [
        "paper-a", "paper-b",       # pilot-1
        "paper-a", "paper-b",       # pilot-2, same two documents
        "paper-a", "paper-b", "paper-c",  # batch, every document
    ]
    assert [guidance for _, guidance in calls[:2]] == ["", ""]
    assert all(guidance.startswith(iterative_eval.GUIDANCE_HEADER) for _, guidance in calls[2:])
    assert (tmp_path / "pipe" / "rounds" / "pilot_2" / "guidance.json").is_file()
    assert (tmp_path / "pipe" / "rounds" / "batch" / "metrics.json").is_file()
    batch_metrics = json.loads((tmp_path / "pipe" / "rounds" / "batch" / "metrics.json").read_text())
    assert batch_metrics["anchor_coverage"]["eligible_link_rate"] == 0.0
    assert batch_metrics["reviewer_effort"]["effort"] >= 1
    assert (tmp_path / "pipe" / "metrics.csv").is_file()
    assert (tmp_path / "pipe" / "metrics.xlsx").is_file()
    rows = list(csv.DictReader((tmp_path / "pipe" / "metrics.csv").open()))
    assert [row["Round"] for row in rows] == ["PILOT_1", "PILOT_2", "BATCH"]
    assert json.loads((tmp_path / "pipe" / "manifest.json").read_text())["pilot_documents"] == ["paper-a", "paper-b"]

    iterative_eval.run_pipeline(config, tmp_path)  # every completed round cell is reused, not re-extracted
    assert len(calls) == 7
